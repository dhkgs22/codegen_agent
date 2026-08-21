"""生成答辩 PPT 的 .pptx 成品文件（深色科技风，7 页）。

运行：python scripts/make_pptx.py
输出：output/slides/AI_Agent_Project_I_答辩.pptx
"""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# 配色（深色科技风）
BG = RGBColor(26, 26, 46)          # #1A1A2E
ACCENT = RGBColor(74, 144, 217)    # #4A90D9 亮蓝
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(180, 185, 200)
FONT = "微软雅黑"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def _set_bg(slide):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG


def _add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = ACCENT
    run.font.name = FONT
    return box


def _add_body(slide, lines, top=1.7, size=18):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(5.3))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ""
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = WHITE if not line.startswith("  ") else GRAY
        run.font.name = FONT
        p.space_after = Pt(12)
    return box


def _add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _blank_slide():
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局
    _set_bg(slide)
    return slide


# ---------- 第 1 页：封面 ----------
s = _blank_slide()
_add_title(s, "多智能体代码助手")
sub = s.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(11.7), Inches(2.0))
tf = sub.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = "项目 I · 2026 暑期 AI Agent 成长实践"
r.font.size = Pt(24); r.font.color.rgb = WHITE; r.font.name = FONT
p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "汇报人：小明"
r2.font.size = Pt(18); r2.font.color.rgb = GRAY; r2.font.name = FONT
_add_notes(s, "开场：老师们好，我的项目是多智能体代码助手，用 AI 自动完成写代码、运行、修Bug、写文档的整个闭环。")

# ---------- 第 2 页：项目背景与目标 ----------
s = _blank_slide()
_add_title(s, "项目背景与目标")
_add_body(s, [
    "问题：日常开发中「写代码、跑代码、改Bug、写文档」占用大量时间",
    "思路：让多个 AI Agent 各司其职，协作完成这条流水线",
    "目标1：输入一句自然语言，自动生成可运行的 Python 程序",
    "目标2：运行失败时，Debug Agent 自动定位并修复（至多 N 轮）",
    "目标3：成功后自动生成 README 文档，全流程可视化",
])
_add_notes(s, "这一页讲清楚为什么做：把写代码中重复、机械的部分交给多Agent协作完成。")

# ---------- 第 3 页：系统架构 ----------
s = _blank_slide()
_add_title(s, "系统架构：多智能体编排")
_add_body(s, [
    "Coordinator（总控编排）",
    "  ├── Code Agent   : 自然语言 → 生成 Python 代码",
    "  ├── Debug Agent  : 沙箱运行 → 失败自动修复（循环）",
    "  ├── Doc Agent    : 运行通过 → 生成 README.md",
    "  └── 事件流(EventBus) : 全过程实时可见，前端/CLI 渲染",
])
_add_notes(s, "核心创新是多Agent分工：每个Agent职责单一，通过Coordinator编排和事件流联动。")

# ---------- 第 4 页：核心技术实现 ----------
s = _blank_slide()
_add_title(s, "核心技术实现")
_add_body(s, [
    "LLM 抽象层：默认 DeepSeek，接口统一，可换 Claude",
    "沙箱执行：subprocess 隔离 + 临时目录 + 超时保护",
    "修复循环：失败 → 报错喂回 LLM → 重跑，至多 N 轮",
    "事件流：结构化事件经 EventBus 广播，CLI/SSE/Web 三种消费端",
    "安全：API Key 存 .env 不入库，沙箱不装未知依赖",
])
_add_notes(s, "重点讲沙箱和安全设计，体现工程素养。")

# ---------- 第 5 页：功能演示 ----------
s = _blank_slide()
_add_title(s, "功能演示")
_add_body(s, [
    "命令行：python -m app.cli \"写一个计算斐波那契数列的程序\"",
    "Web 界面：Streamlit 实时观看 Agent 工作过程",
    "API：POST /api/run（同步）+ GET /api/run/stream（SSE 流式）",
    "真实效果：生成→运行→修Bug→README，端到端 7.5 秒完成",
])
_add_notes(s, "这里切到演示视频或现场demo，展示真实运行。")

# ---------- 第 6 页：成果与数据 ----------
s = _blank_slide()
_add_title(s, "成果与数据")
_add_body(s, [
    "测试：15 个 pytest 测试全部通过（Fake LLM，不消耗 token）",
    "CI/CD：GitHub Actions 自动验证，推送即跑，全绿",
    "版本：GitHub Release v1.0.0，含完整 Release Note",
    "工程规范：MIT License + CONTRIBUTING 贡献指南",
    "三种入口：CLI / Web / API，全部实测可用",
])
_add_notes(s, "用数据说话，展示仓库地址和CI截图。")

# ---------- 第 7 页：总结与展望 ----------
s = _blank_slide()
_add_title(s, "总结与展望")
_add_body(s, [
    "收获：掌握多Agent编排、事件驱动、沙箱隔离、CI/CD 全套工程方法",
    "不足：仅支持单文件 Python 项目，Debug 依赖 LLM 质量",
    "展望：支持多文件与依赖安装、加入测试生成 Agent、Benchmark 统计",
    "",
    "感谢聆听，敬请指正！",
])
_add_notes(s, "收尾：感谢老师，欢迎提问。")

out = "output/slides/AI_Agent_Project_I_答辩.pptx"
prs.save(out)
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except (AttributeError, ValueError):
        pass
print(f"OK: {out} ({len(prs.slides)} slides)")
